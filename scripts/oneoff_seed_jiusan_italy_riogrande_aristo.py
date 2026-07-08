"""一次性脚本:为九三大豆补录意大利/里奥格兰德/阿里斯托历史 release_batch。

本脚本只建立 release_batches 批次锚点,不回灌 wagon/container 明细。

口径来源:
- 用户 2026-07-08 明确确认:
  - 意大利:仅 1 个集装箱批次
  - 里奥格兰德: lot01 集装箱, lot02 散粮车 15000 吨
  - 阿里斯托:仅 1 个集装箱批次
- Excel 文件只用于提取首末发运日期与历史计数口径
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
from sop_hub.utils.time import now_iso_beijing  # noqa: E402

SOP_DB = REPO / "data" / "sop_agent.db"
DOWNLOADS = Path("/Users/qicai21/Downloads")


@dataclass
class ManifestSummary:
    first_date: str
    last_date: str
    row_count: int
    wagon_count: int


def stable_hash(*parts: str) -> str:
    return hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()


def _norm_date(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return None
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _collect_manifest_summary(path: Path, sheet_names: list[str], *, is_container: bool) -> ManifestSummary:
    wb = load_workbook(path, read_only=True, data_only=True)
    dates: list[str] = []
    row_count = 0
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            vals = list(row)
            if not vals:
                continue
            try:
                int(vals[0])
            except Exception:
                continue
            d = _norm_date(vals[1] if len(vals) > 1 else None)
            if d:
                dates.append(d)
            row_count += 1
    if not dates:
        raise RuntimeError(f"{path.name} / {sheet_names} 未提取到有效日期")
    wagon_count = row_count // 2 if is_container else row_count
    return ManifestSummary(
        first_date=min(dates),
        last_date=max(dates),
        row_count=row_count,
        wagon_count=wagon_count,
    )


def build_batches(now: str) -> list[dict]:
    italy = _collect_manifest_summary(DOWNLOADS / "意大利.xlsx", ["新台子"], is_container=True)
    rio_c = _collect_manifest_summary(DOWNLOADS / "里奥格兰德.xlsx", ["新台子"], is_container=True)
    rio_b = _collect_manifest_summary(DOWNLOADS / "里奥格兰德.xlsx", ["新台子-散粮车"], is_container=False)
    aristo_wb = load_workbook(DOWNLOADS / "阿里斯托.xlsx", read_only=True, data_only=True)
    aristo_sheets = [s for s in aristo_wb.sheetnames if s != "发运进度"]
    aristo = _collect_manifest_summary(DOWNLOADS / "阿里斯托.xlsx", aristo_sheets, is_container=True)

    common = {
        "project": "jiusan",
        "contract_no": "JGWL-JZTS-DD-202601",
        "cargo_name": "大豆",
        "consignor": "锦州港物流发展有限公司",
        "origin_station": "高桥镇",
        "trade_type": "外贸进口",
        "destination_station": "新台子",
        "remaining_quantity": 0.0,
        "remaining_weight_tons": 0.0,
        "shipped_weight_tons": 0.0,
        "dispatch_status": "confirmed_received",
        "dispatch_status_note": "historical_manifest_seed_20260708",
        "dispatch_status_updated_at": now,
    }
    rows = [
        {
            "ship_name_cn": "意大利",
            "ship_name_en": "MV BULK ITALY",
            "voyage_no": "2525",
            "cargo_origin": "阿根廷",
            "arrived_at": "2025-11-07 06:30",
            "departed_at": "2025-11-14 19:30",
            "manifest_total_tons": 68562.20,
            "customs_release_tons": 68562.20,
            "ship_name": "意大利",
            "batch_sequence": "lot01",
            "transport_mode": "铁路集装箱",
            "yard_location": "国家粮食和物资储备局辽宁局三三0处专用线",
            "consignee": "国家粮食和物资储备局辽宁局三三0处",
            "notice_date": italy.first_date,
            "batch_date": italy.first_date,
            "batch_quantity": 68562.20,
            "total_planned_quantity": 68562.20,
            "batch_count": italy.wagon_count,
            "actual_wagon_count": italy.wagon_count,
            "confirmed_received_at": italy.last_date + "T20:00:00+08:00",
            "cargo_product_name": "大豆-集装箱",
            "cargo_name_detail": "大豆-集装箱",
            "source_file_name": "意大利.xlsx",
            "searchable_text": "九三大豆 意大利 lot01 铁路集装箱 高桥镇 新台子 三三0处",
            "import_ship_name": "MV BULK ITALY",
            "source_meta": {
                "src": "historical_manifest_seed_20260708",
                "sheet_names": ["新台子"],
                "first_event_date": italy.first_date,
                "last_event_date": italy.last_date,
                "manifest_row_count": italy.row_count,
                "manifest_wagon_count": italy.wagon_count,
            },
        },
        {
            "ship_name_cn": "里奥格兰德",
            "ship_name_en": "MV GREEN RIO GRANDE",
            "voyage_no": "2510",
            "cargo_origin": "巴西",
            "arrived_at": "2025-11-28 15:18",
            "departed_at": "2025-12-12 23:30",
            "manifest_total_tons": 63474.44,
            "customs_release_tons": 63474.44,
            "ship_name": "里奥格兰德",
            "batch_sequence": "lot01",
            "transport_mode": "铁路集装箱",
            "yard_location": "国家粮食和物资储备局辽宁局三三0处专用线",
            "consignee": "国家粮食和物资储备局辽宁局三三0处",
            "notice_date": rio_c.first_date,
            "batch_date": rio_c.first_date,
            "batch_quantity": 48474.44,
            "total_planned_quantity": 48474.44,
            "batch_count": rio_c.wagon_count,
            "actual_wagon_count": rio_c.wagon_count,
            "confirmed_received_at": rio_c.last_date + "T20:00:00+08:00",
            "cargo_product_name": "大豆-集装箱",
            "cargo_name_detail": "大豆-集装箱",
            "source_file_name": "里奥格兰德.xlsx",
            "searchable_text": "九三大豆 里奥格兰德 lot01 铁路集装箱 高桥镇 新台子 三三0处",
            "import_ship_name": "MV GREEN RIO GRANDE",
            "source_meta": {
                "src": "historical_manifest_seed_20260708",
                "sheet_names": ["新台子"],
                "first_event_date": rio_c.first_date,
                "last_event_date": rio_c.last_date,
                "manifest_row_count": rio_c.row_count,
                "manifest_wagon_count": rio_c.wagon_count,
                "lot_planned_tons": 48474.44,
                "split_basis": "用户确认: 全船63474.44吨, lot02散粮15000吨, lot01集装箱=48474.44吨",
            },
        },
        {
            "ship_name_cn": "里奥格兰德",
            "ship_name_en": "MV GREEN RIO GRANDE",
            "voyage_no": "2510",
            "cargo_origin": "巴西",
            "arrived_at": "2025-11-28 15:18",
            "departed_at": "2025-12-12 23:30",
            "manifest_total_tons": 63474.44,
            "customs_release_tons": 63474.44,
            "ship_name": "里奥格兰德",
            "batch_sequence": "lot02",
            "transport_mode": "铁路散粮车",
            "yard_location": "九三集团铁岭大豆科技有限公司专用线",
            "consignee": "九三集团铁岭大豆科技有限公司",
            "notice_date": rio_b.first_date,
            "batch_date": rio_b.first_date,
            "batch_quantity": 15000.0,
            "total_planned_quantity": 15000.0,
            "batch_count": rio_b.wagon_count,
            "actual_wagon_count": rio_b.wagon_count,
            "confirmed_received_at": rio_b.last_date + "T20:00:00+08:00",
            "cargo_product_name": "大豆-散粮车",
            "cargo_name_detail": "大豆-散粮车",
            "source_file_name": "里奥格兰德.xlsx",
            "searchable_text": "九三大豆 里奥格兰德 lot02 铁路散粮车 高桥镇 新台子 九三专用线",
            "import_ship_name": "MV GREEN RIO GRANDE",
            "source_meta": {
                "src": "historical_manifest_seed_20260708",
                "sheet_names": ["新台子-散粮车"],
                "first_event_date": rio_b.first_date,
                "last_event_date": rio_b.last_date,
                "manifest_row_count": rio_b.row_count,
                "manifest_wagon_count": rio_b.wagon_count,
                "lot_planned_tons": 15000.0,
                "split_basis": "用户确认: lot02散粮车15000吨",
            },
        },
        {
            "ship_name_cn": "阿里斯托",
            "ship_name_en": "MV ARISTO OCEAN",
            "voyage_no": "2520",
            "cargo_origin": "阿根廷/乌拉圭",
            "arrived_at": "2025-11-18 06:54",
            "departed_at": "2025-11-28 01:18",
            "manifest_total_tons": 68651.54,
            "customs_release_tons": 68651.54,
            "ship_name": "阿里斯托",
            "batch_sequence": "lot01",
            "transport_mode": "铁路集装箱",
            "yard_location": "国家粮食和物资储备局辽宁局三三0处专用线",
            "consignee": "国家粮食和物资储备局辽宁局三三0处",
            "notice_date": aristo.first_date,
            "batch_date": aristo.first_date,
            "batch_quantity": 68651.54,
            "total_planned_quantity": 68651.54,
            "batch_count": aristo.wagon_count,
            "actual_wagon_count": aristo.wagon_count,
            "confirmed_received_at": aristo.last_date + "T20:00:00+08:00",
            "cargo_product_name": "大豆-集装箱",
            "cargo_name_detail": "大豆-集装箱",
            "source_file_name": "阿里斯托.xlsx",
            "searchable_text": "九三大豆 阿里斯托 lot01 铁路集装箱 高桥镇 新台子 三三0处",
            "import_ship_name": "MV ARISTO OCEAN",
            "source_meta": {
                "src": "historical_manifest_seed_20260708",
                "sheet_names": aristo_sheets,
                "first_event_date": aristo.first_date,
                "last_event_date": aristo.last_date,
                "manifest_row_count": aristo.row_count,
                "manifest_wagon_count": aristo.wagon_count,
            },
        },
    ]

    out: list[dict] = []
    for row in rows:
        ship_name = row["ship_name"]
        lot = row["batch_sequence"]
        notice_date = row["notice_date"]
        batch_key = f"jiusan|{ship_name}|{lot}|{notice_date}"
        source_json = {
            "ship_name_cn": row["ship_name_cn"],
            "ship_name_en": row["ship_name_en"],
            "voyage_no": row["voyage_no"],
            "cargo_origin": row["cargo_origin"],
            "arrived_at": row["arrived_at"],
            "departed_at": row["departed_at"],
            "manifest_total_tons": row["manifest_total_tons"],
            "customs_release_tons": row["customs_release_tons"],
            **row["source_meta"],
        }
        out.append({
            **common,
            "id": stable_hash("jiusan", ship_name, lot, notice_date),
            "batch_key": batch_key,
            "ship_name": ship_name,
            "batch_sequence": lot,
            "transport_mode": row["transport_mode"],
            "destination_station": "新台子",
            "yard_location": row["yard_location"],
            "consignee": row["consignee"],
            "notice_date": notice_date,
            "batch_date": row["batch_date"],
            "batch_quantity": row["batch_quantity"],
            "total_planned_quantity": row["total_planned_quantity"],
            "batch_count": row["batch_count"],
            "actual_wagon_count": row["actual_wagon_count"],
            "confirmed_received_at": row["confirmed_received_at"],
            "cargo_product_name": row["cargo_product_name"],
            "cargo_name_detail": row["cargo_name_detail"],
            "source_file_name": row["source_file_name"],
            "source_json": json.dumps(source_json, ensure_ascii=False),
            "searchable_text": row["searchable_text"],
            "import_ship_name": row["import_ship_name"],
            "customs_release_qty": row["customs_release_tons"],
            "quantity_tons": row["batch_quantity"],
            "created_at": now,
            "updated_at": now,
        })
    return out


def main(apply: bool = False):
    now = now_iso_beijing()
    conn = sqlite3.connect(str(SOP_DB))
    batches = build_batches(now)
    cur = conn.cursor()
    print(f"准备插 {len(batches)} 个 release_batch:")
    for b in batches:
        print(
            f"  {b['id'][:12]}  {b['ship_name']:6s} {b['batch_sequence']} "
            f"{b['transport_mode']:10s} qty={b['batch_quantity']:.2f}t "
            f"date={b['batch_date']} status={b['dispatch_status']}"
        )
        cols = list(b.keys())
        placeholders = ",".join(["?"] * len(cols))
        cur.execute(
            f"INSERT OR REPLACE INTO release_batches ({','.join(cols)}) VALUES ({placeholders})",
            [b[c] for c in cols],
        )
    if apply:
        conn.commit()
        print("\nCOMMIT ✓")
    else:
        conn.rollback()
        print("\nDRY-RUN(加 --apply 真正提交)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    main(apply=args.apply)
