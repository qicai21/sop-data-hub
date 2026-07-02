from __future__ import annotations

import hashlib
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
try:
    from docx import Document
except ModuleNotFoundError:  # pragma: no cover - optional at import time
    Document = None


REPO_ROOT = Path(__file__).resolve().parents[3]
JIUSAN_YAML = REPO_ROOT / "config" / "project_sops" / "jiusan.yaml"
DOC_ROOT = REPO_ROOT / "reports" / "fee_docs" / "jiusan"

_TRACK_PATTERN = re.compile(r"(煤[一二三四五六七八九])|([一二三四五六七八九十\d]+道)")


def stable_hash(*parts: Any) -> str:
    return hashlib.sha1("|".join(str(p) for p in parts).encode("utf-8")).hexdigest()[:24]


def load_route_c_config() -> dict[str, Any]:
    data = yaml.safe_load(JIUSAN_YAML.read_text(encoding="utf-8")) or {}
    return ((data.get("cost_structure") or {}).get("route_c_bulk") or {})


def is_real_track(track: str | None) -> bool:
    return bool(track and _TRACK_PATTERN.search(track))


def display_track(track: str | None) -> str:
    if is_real_track(track):
        return str(track).strip()
    return "待补录"


def build_docx_path(notice_date: str, ship_name: str, track: str | None, car_count: int) -> Path:
    month = notice_date[:7]
    safe_track = display_track(track).replace("/", "_")
    name = f"{notice_date}_{ship_name}_{safe_track}_{car_count}车_现场确认单.docx"
    return DOC_ROOT / month / name


def freight_fee_yuan(raw_fee: Any) -> float:
    try:
        return round(float(raw_fee) / 100.0, 2)
    except (TypeError, ValueError):
        return 0.0


def billing_weight(row: dict[str, Any]) -> float:
    v = row.get("marked_weight")
    if v not in (None, ""):
        try:
            return float(v)
        except (TypeError, ValueError):
            pass
    v = row.get("computed_loading_weight")
    if v not in (None, ""):
        try:
            return float(v)
        except (TypeError, ValueError):
            pass
    model = str(row.get("car_model") or "")
    if model.startswith("L70"):
        return 69.0
    return 61.0


def load_contract_terms(conn: sqlite3.Connection, project_id: str, route_code: str) -> dict[str, dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT fee_code, fee_name, charge_side, pricing_basis, pricing_unit, default_rate,
               tax_rate, counterparty, settle_party, note
        FROM contract_fee_terms
        WHERE project_id=? AND route_code=? AND enabled=1
        """,
        (project_id, route_code),
    ).fetchall()
    return {str(r["fee_code"]): dict(r) for r in rows}


def load_line_rates(
    conn: sqlite3.Connection,
    project_id: str,
    route_code: str,
    fee_code: str,
) -> dict[str, float]:
    rows = conn.execute(
        """
        SELECT line_name, rate
        FROM contract_line_rates
        WHERE project_id=? AND route_code=? AND fee_code=? AND enabled=1
        """,
        (project_id, route_code, fee_code),
    ).fetchall()
    return {str(r["line_name"]).strip(): float(r["rate"]) for r in rows}


def load_weight_confirmation(
    conn: sqlite3.Connection,
    release_batch_id: str,
    route_code: str,
    weight_type: str = "port_weighing_departure",
) -> float | None:
    row = conn.execute(
        """
        SELECT confirmed_weight
        FROM shipment_weight_confirmation
        WHERE release_batch_id=? AND route_code=? AND weight_type=?
        """,
        (release_batch_id, route_code, weight_type),
    ).fetchone()
    if not row:
        return None
    try:
        return float(row["confirmed_weight"])
    except (TypeError, ValueError):
        return None


def allocated_confirmed_weight(
    total_confirmed_weight: float | None,
    batch_marked_weight: float,
    ship_marked_weight: float,
) -> float:
    if total_confirmed_weight in (None, ""):
        return 0.0
    total = float(total_confirmed_weight or 0.0)
    batch = float(batch_marked_weight or 0.0)
    ship = float(ship_marked_weight or 0.0)
    if total <= 0 or batch <= 0 or ship <= 0:
        return 0.0
    return round(total * batch / ship, 2)


def resolve_line_rate(
    line_name: str | None,
    rate_map: dict[str, float],
    default_rate: float | None = None,
) -> tuple[float | None, str]:
    name = str(line_name or "").strip()
    if name and name in rate_map:
        return rate_map[name], name
    return default_rate, name


@dataclass
class FeeItemCalc:
    code: str
    name: str
    side: str
    qty: float
    qty_unit: str
    amount: float
    price: float | None
    settle_party: str
    counterparty: str
    pricing_basis: str
    pricing_basis_value: float
    tax_rate: float
    document_flow_type: str
    source_ref: str
    service_no: str = ""


def calc_fee_items(
    *,
    terms: dict[str, dict[str, Any]],
    total_weight: float,
    freight_sum_yuan: float,
    car_count: int,
    line_rate: float | None,
    box_trip_count: int = 0,
    code_weights: dict[str, float] | None = None,
    source_ref: str,
) -> list[FeeItemCalc]:
    out: list[FeeItemCalc] = []
    overrides = code_weights or {}
    for code, item in terms.items():
        name = str(item["fee_name"])
        side = str(item["charge_side"])
        basis = str(item.get("pricing_basis") or "")
        tax = float(item.get("tax_rate", 0) or 0)
        settle_party = str(item.get("settle_party") or "")
        counterparty = str(item.get("counterparty") or "")
        doc_type = "onsite_confirm_sheet" if code in {"aux_bulk_loading", "aux_bulk_inspection"} else ""
        service_no = "19" if code == "aux_bulk_loading" else "20" if code == "aux_bulk_inspection" else ""
        rate = float(item.get("default_rate", 0) or 0)

        if basis == "freight_fee_sum":
            qty = round(total_weight, 2)
            amount = round(freight_sum_yuan, 2)
            price = round(amount / qty, 4) if qty else 0.0
            qty_unit = "ton"
            pricing_basis = "actual_freight_fee"
            pricing_basis_value = round(freight_sum_yuan, 2)
        elif basis in {"confirmed_weight", "marked_weight", "marked_weight_tiered"}:
            weight = float(overrides.get(code, total_weight))
            qty = round(weight, 2)
            amount = round(rate * qty, 2)
            price = rate
            qty_unit = "ton"
            pricing_basis = basis
            pricing_basis_value = round(weight, 2)
        elif basis == "marked_weight_x_line_rate":
            if line_rate is None:
                continue
            weight = float(overrides.get(code, total_weight))
            qty = round(weight, 2)
            amount = round(float(line_rate) * qty, 2)
            price = float(line_rate)
            qty_unit = "ton"
            pricing_basis = basis
            pricing_basis_value = round(weight, 2)
        elif basis == "car_count":
            qty = float(car_count)
            amount = round(rate * qty, 2)
            price = rate
            qty_unit = "car"
            pricing_basis = basis
            pricing_basis_value = float(car_count)
        elif basis == "box_trip_count":
            qty = float(box_trip_count)
            amount = round(rate * qty, 2)
            price = rate
            qty_unit = "box"
            pricing_basis = basis
            pricing_basis_value = float(box_trip_count)
        else:
            continue

        out.append(
            FeeItemCalc(
                code=code,
                name=name,
                side=side,
                qty=qty,
                qty_unit=qty_unit,
                amount=amount,
                price=price,
                settle_party=settle_party,
                counterparty=counterparty,
                pricing_basis=pricing_basis,
                pricing_basis_value=pricing_basis_value,
                tax_rate=tax,
                document_flow_type=doc_type or "",
                source_ref=source_ref,
                service_no=service_no,
            )
        )
    return out


def render_onsite_confirm_docx(
    *,
    path: Path,
    notice_date: str,
    project_name: str,
    track: str | None,
    car_count: int,
    car_nos: list[str],
    fee_items: list[FeeItemCalc],
    ship_name: str,
) -> None:
    if Document is None:
        raise ModuleNotFoundError("python-docx is required to render onsite confirm docs")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("锦州港装卸辅助作业现场确认单", level=1)
    doc.add_paragraph("-----------------------")
    doc.add_paragraph(f"发生日期: {notice_date}")
    doc.add_paragraph(f"项目名称: {project_name}")
    doc.add_paragraph(f"船名/批次: {ship_name}")
    doc.add_paragraph(f"道线场地: {display_track(track)}")
    if not is_real_track(track):
        doc.add_paragraph("备注: 当前台账来源为按船发运报表,真实作业股道待补录。")

    doc.add_paragraph("作业项目:")
    for item in fee_items:
        if item.code == "aux_bulk_loading":
            doc.add_paragraph(
                f"  -[ ] #{item.service_no} {item.name}/特殊平整、清扫归集及皮带机配合"
            )
        elif item.code == "aux_bulk_inspection":
            doc.add_paragraph(
                f"  -[ ] #{item.service_no} {item.name}/铁路及海关监管"
            )

    doc.add_paragraph(f"数量: {car_count} 车")
    doc.add_paragraph(f"车号/箱号: {', '.join(car_nos)}")
    doc.add_paragraph("")
    doc.add_paragraph("作业单位/班组:")
    doc.add_paragraph("")
    doc.add_paragraph("委托方(签字):")
    doc.add_paragraph("")
    doc.add_paragraph("作业方(签字):")
    doc.save(path)
