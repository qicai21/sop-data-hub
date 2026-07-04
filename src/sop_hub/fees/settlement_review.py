from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from typing import Iterable, Sequence


SETTLEMENT_STATUSES = ("结算", "挂起", "排除")
REVIEW_SHEET_NAME = "货运记录表"


@dataclass(frozen=True)
class ShipmentReviewRow:
    row_id: str
    project_id: str
    ship_name: str
    lot: str
    transport_mode: str
    shipment_date: date | str
    track: str = ""
    route: str = ""
    source_table: str = ""
    source_ids: str = ""
    wagon_count: int = 0
    container_count: int = 0
    quantity: Decimal | int | float | str = Decimal("0")
    quantity_unit: str = ""
    settlement_status: str = "结算"
    merge_key: str = ""
    quantity_adjustment: Decimal | int | float | str = Decimal("0")
    info_summary: str = ""
    remarks: str = ""

    def __post_init__(self) -> None:
        if self.settlement_status not in SETTLEMENT_STATUSES:
            raise ValueError(f"invalid settlement_status: {self.settlement_status}")

    @property
    def effective_quantity(self) -> Decimal:
        return _to_decimal(self.quantity) + _to_decimal(self.quantity_adjustment)

    @property
    def settlement_key(self) -> str:
        return self.merge_key.strip() or self.row_id


@dataclass(frozen=True)
class SettlementUnit:
    settlement_key: str
    rows: tuple[ShipmentReviewRow, ...] = field(default_factory=tuple)

    @property
    def quantity(self) -> Decimal:
        return sum((row.effective_quantity for row in self.rows), Decimal("0"))

    @property
    def wagon_count(self) -> int:
        return sum(row.wagon_count for row in self.rows)

    @property
    def container_count(self) -> int:
        return sum(row.container_count for row in self.rows)

    @property
    def quantity_unit(self) -> str:
        units = [row.quantity_unit for row in self.rows if row.quantity_unit]
        return units[0] if units else ""

    @property
    def transport_mode(self) -> str:
        modes = [row.transport_mode for row in self.rows if row.transport_mode]
        return modes[0] if modes else ""


REVIEW_HEADERS = [
    "序号",
    "项目",
    "船名",
    "lot",
    "运输方式",
    "日期",
    "道线",
    "路线",
    "来源表",
    "来源记录",
    "车数",
    "箱数",
    "原始数量",
    "计量单位",
    "结算状态",
    "合并标识",
    "调增减数量",
    "有效数量",
    "信息摘要",
    "备注",
]


def build_settlement_units(rows: Iterable[ShipmentReviewRow]) -> list[SettlementUnit]:
    grouped: dict[str, list[ShipmentReviewRow]] = {}
    for row in rows:
        if row.settlement_status != "结算":
            continue
        grouped.setdefault(row.settlement_key, []).append(row)
    return [
        SettlementUnit(settlement_key=key, rows=tuple(group_rows))
        for key, group_rows in grouped.items()
    ]


def add_review_sheet(workbook, rows: Sequence[ShipmentReviewRow], sheet_name: str = REVIEW_SHEET_NAME):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    ws = workbook.create_sheet(sheet_name, 0)
    ws.append(REVIEW_HEADERS)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, row in enumerate(rows, start=1):
        ws.append(
            [
                index,
                row.project_id,
                row.ship_name,
                row.lot,
                row.transport_mode,
                str(row.shipment_date),
                row.track,
                row.route,
                row.source_table,
                row.source_ids,
                row.wagon_count,
                row.container_count,
                _excel_number(row.quantity),
                row.quantity_unit,
                row.settlement_status,
                row.merge_key,
                _excel_number(row.quantity_adjustment),
                _excel_number(row.effective_quantity),
                row.info_summary,
                row.remarks,
            ]
        )

    status_col = REVIEW_HEADERS.index("结算状态") + 1
    status_letter = get_column_letter(status_col)
    last_row = max(2, len(rows) + 1)
    validation = DataValidation(
        type="list",
        formula1=f'"{",".join(SETTLEMENT_STATUSES)}"',
        allow_blank=False,
    )
    ws.add_data_validation(validation)
    validation.add(f"{status_letter}2:{status_letter}{last_row}")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_HEADERS))}{last_row}"
    widths = {
        "A": 8,
        "B": 14,
        "C": 14,
        "D": 10,
        "E": 12,
        "F": 14,
        "G": 10,
        "H": 18,
        "I": 24,
        "J": 24,
        "O": 12,
        "P": 16,
        "Q": 14,
        "R": 12,
        "S": 36,
        "T": 24,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    return ws


def _to_decimal(value: Decimal | int | float | str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value or "0"))


def _excel_number(value: Decimal | int | float | str) -> int | float:
    decimal_value = _to_decimal(value)
    if decimal_value == decimal_value.to_integral_value():
        return int(decimal_value)
    return float(decimal_value)


__all__ = [
    "REVIEW_HEADERS",
    "REVIEW_SHEET_NAME",
    "SETTLEMENT_STATUSES",
    "SettlementUnit",
    "ShipmentReviewRow",
    "add_review_sheet",
    "build_settlement_units",
]
