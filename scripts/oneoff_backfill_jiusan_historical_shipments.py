"""一次性脚本: 九三历史 95306 明细回灌 + 当前 container hph 回填。

用途:
1. 将意大利 / 里奥格兰德 / 阿里斯托历史批次,按 Excel 货票号去 95306 拉明细,
   回灌到 wagon_container_shipments / wagon_shipments。
2. 回填九三现有 wagon 明细里的 hph / waybill_no,便于后续分票与核对。

口径:
- 单一事实源仍是 95306;Excel 只负责给出需要纳入的货票号集合。
- 95306 缺失的历史票(如意大利 2025-12 及 2026-01-01 早段缺口)不造数。
- 只处理新台子方向(历史补录目标批次均为新台子)。

用法:
  python3 scripts/oneoff_backfill_jiusan_historical_shipments.py
  python3 scripts/oneoff_backfill_jiusan_historical_shipments.py --apply
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from sop_hub.sop.shipped_weight import compute_for_release_batch  # noqa: E402
from sop_hub.utils.time import now_iso_beijing  # noqa: E402

SOP_DB = REPO / "data" / "sop_agent.db"
RAIL_DB = Path("/Users/qicai21/projects/repos/rail95306-sync/runtime/95306_collection.sqlite3")
DOWNLOADS = Path("/Users/qicai21/Downloads")
PROJECT = "jiusan"


@dataclass(frozen=True)
class BatchSpec:
    ship_name: str
    batch_sequence: str
    transport_type: str  # container / bulk
    workbook: str
    sheets: tuple[str, ...]
    format_type: str  # standard_waybill / waybill_in_box_rows


SPECS: tuple[BatchSpec, ...] = (
    BatchSpec("意大利", "lot01", "container", "意大利.xlsx", ("新台子",), "standard_waybill"),
    BatchSpec("里奥格兰德", "lot01", "container", "里奥格兰德.xlsx", ("新台子",), "standard_waybill"),
    BatchSpec("里奥格兰德", "lot02", "bulk", "里奥格兰德.xlsx", ("新台子-散粮车",), "standard_waybill"),
    BatchSpec(
        "阿里斯托",
        "lot01",
        "container",
        "阿里斯托.xlsx",
        (
            "1.25", "1.27", "1.28", "1.29", "1.30", "1.31", "2.1", "2.2", "2.3",
            "2.4", "2.5", "2.6", "2.8", "2.10", "2.24", "2.28", "3.2", "3.4",
            "3.6", "3.8", "3.10", "3.12", "3.14", "3.16", "3.18",
        ),
        "waybill_in_box_rows",
    ),
)


def stable_hash(*parts: str) -> str:
    return hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()


def _parse_hph(raw: str | None) -> str | None:
    if not raw:
        return None
    hits = re.findall(r"GZD[A-Z]{2}[0-9]{6,}", raw)
    return hits[0] if hits else None


def _parse_boxes(raw: str | None) -> list[str]:
    if not raw:
        return []
    try:
        vals = json.loads(raw)
    except Exception:
        return []
    boxes: list[str] = []
    if isinstance(vals, list):
        for val in vals:
            s = str(val or "").strip()
            if s:
                boxes.append(s)
    return boxes


def load_excel_hphs(spec: BatchSpec) -> list[str]:
    path = DOWNLOADS / spec.workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    hphs: list[str] = []
    if spec.format_type == "standard_waybill":
        for sheet in spec.sheets:
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=4, values_only=True):
                if not row or row[0] in (None, "", "合计"):
                    continue
                try:
                    int(row[0])
                except Exception:
                    continue
                hph = str(row[3] or "").strip()
                if hph:
                    hphs.append(hph)
    elif spec.format_type == "waybill_in_box_rows":
        for sheet in spec.sheets:
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=3, values_only=True):
                if not row or row[0] in (None, ""):
                    continue
                try:
                    int(row[0])
                except Exception:
                    continue
                hph = str(row[8] or "").strip() if len(row) > 8 else ""
                if hph:
                    hphs.append(hph)
    else:
        raise RuntimeError(f"unknown format_type={spec.format_type}")
    # 去重并保持顺序
    return list(dict.fromkeys(hphs))


def load_rail_rows() -> tuple[dict[str, dict], dict[str, str]]:
    conn = sqlite3.connect(str(RAIL_DB))
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        """
        SELECT
          ydid, czydid, car_no, car_model, container_numbers_json,
          marked_weight, status_name, latest_stage_key, latest_stage_name,
          latest_event_time, accepted_at, loaded_at, ticketed_at,
          departed_at, arrived_at, delivered_at,
          transport_mode_code, transport_mode_name,
          origin_name, destination_name, cargo_name,
          freight_fee, raw_core_json
        FROM shipments
        WHERE cargo_name LIKE '%豆%'
          AND origin_name='高桥镇'
          AND destination_name='新台子'
          AND ticketed_at >= '2026-01-01'
        ORDER BY ticketed_at, car_no
        """
    ).fetchall()
    conn.close()

    by_hph: dict[str, dict] = {}
    ydid_to_hph: dict[str, str] = {}
    for row in rows:
        item = dict(row)
        hph = _parse_hph(item.get("raw_core_json"))
        if not hph:
            continue
        item["hph"] = hph
        item["boxes"] = _parse_boxes(item.get("container_numbers_json"))
        by_hph[hph] = item
        ydid_to_hph[item["ydid"]] = hph
    return by_hph, ydid_to_hph


def map_dispatch_status(status_name: str) -> str:
    if status_name and "交付" in status_name:
        return "confirmed_received"
    return "loading"


def batch_id_for(conn: sqlite3.Connection, ship_name: str, batch_sequence: str) -> str:
    row = conn.execute(
        """
        SELECT id FROM release_batches
        WHERE project=? AND ship_name=? AND batch_sequence=?
        """,
        (PROJECT, ship_name, batch_sequence),
    ).fetchone()
    if not row:
        raise RuntimeError(f"batch not found: {ship_name} {batch_sequence}")
    return row[0]


def build_container_row(rail_row: dict, box_no: str, pos: int, batch_id: str, ship_name: str, now: str) -> dict:
    return {
        "id": stable_hash(str(rail_row["car_no"]), box_no, str(rail_row["ydid"])),
        "car_no": rail_row["car_no"],
        "box_no": box_no,
        "box_position": pos,
        "ydid": rail_row["ydid"],
        "czydid": rail_row["czydid"],
        "waybill_no": rail_row["hph"],
        "batch_id": batch_id,
        "car_model": rail_row["car_model"] or "",
        "ticketed_at": rail_row["ticketed_at"] or "",
        "departed_at": rail_row["departed_at"] or "",
        "arrived_at": rail_row["arrived_at"] or "",
        "delivered_at": rail_row["delivered_at"] or "",
        "accepted_at": rail_row["accepted_at"] or "",
        "loaded_at": rail_row["loaded_at"] or "",
        "status_name": rail_row["status_name"] or "",
        "latest_stage_key": rail_row["latest_stage_key"] or "",
        "latest_stage_name": rail_row["latest_stage_name"] or "",
        "latest_event_time": rail_row["latest_event_time"] or "",
        "origin_name": rail_row["origin_name"] or "",
        "destination_name": rail_row["destination_name"] or "",
        "transport_mode_code": rail_row["transport_mode_code"] or "",
        "transport_mode_name": rail_row["transport_mode_name"] or "",
        "cargo_name": rail_row["cargo_name"] or "大豆",
        "marked_weight": rail_row["marked_weight"],
        "project_id": PROJECT,
        "ship_name": ship_name,
        "consignor": "锦州港物流发展有限公司",
        "consignee": "国家粮食和物资储备局辽宁局三三0处",
        "dispatch_status": map_dispatch_status(rail_row["status_name"] or ""),
        "source_message_id": f"historical_manifest_backfill_{ship_name}_{batch_id}",
        "source_group_id": "",
        "created_at": now,
        "updated_at": now,
        "cycle_id": None,
        "portal_id": None,
        "portal_uploaded_at": None,
        "hph": rail_row["hph"],
        "reconciled_at": None,
        "reconcile_source_ref": None,
        "freight_fee": rail_row["freight_fee"],
        "detail_json": rail_row.get("raw_core_json") or "",
        "loading_line": None,
        "dispatch_train_code": None,
    }


def build_bulk_row(rail_row: dict, batch_id: str, ship_name: str, now: str) -> dict:
    ds = map_dispatch_status(rail_row["status_name"] or "")
    confirmed_at = rail_row["delivered_at"] if ds == "confirmed_received" else None
    return {
        "id": stable_hash("bulk", str(rail_row["ydid"])),
        "departure_id": None,
        "batch_id": batch_id,
        "car_no": rail_row["car_no"],
        "car_model": rail_row["car_model"] or "",
        "cargo_name": rail_row["cargo_name"] or "大豆",
        "shipper_name": "锦州港物流发展有限公司",
        "consignee_name": "九三集团铁岭大豆科技有限公司",
        "origin_name": rail_row["origin_name"] or "",
        "destination_name": rail_row["destination_name"] or "",
        "ticketed_at": rail_row["ticketed_at"] or "",
        "departed_at": rail_row["departed_at"] or "",
        "arrived_at": rail_row["arrived_at"] or "",
        "status_name": rail_row["status_name"] or "",
        "freight_fee": rail_row["freight_fee"],
        "detail_json": rail_row.get("raw_core_json") or "",
        "created_at": now,
        "updated_at": now,
        "delivered_at": rail_row["delivered_at"] or "",
        "confirmed_received_at": confirmed_at,
        "container_no": None,
        "waybill_no": rail_row["hph"],
        "project_id": PROJECT,
        "ship_name": ship_name,
        "dispatch_status": ds,
        "source_message_id": f"historical_manifest_backfill_{ship_name}_{batch_id}",
        "source_group_id": "",
        "ydid": rail_row["ydid"],
        "czydid": rail_row["czydid"],
        "transport_mode_code": rail_row["transport_mode_code"] or "",
        "transport_mode_name": rail_row["transport_mode_name"] or "",
        "marked_weight": rail_row["marked_weight"],
        "cargo_count": None,
        "container_numbers_json": None,
        "accepted_at": rail_row["accepted_at"] or "",
        "loaded_at": rail_row["loaded_at"] or "",
        "latest_stage_key": rail_row["latest_stage_key"] or "",
        "latest_stage_name": rail_row["latest_stage_name"] or "",
        "latest_event_time": rail_row["latest_event_time"] or "",
        "computed_loading_weight": None,
        "weight_rule_basis": None,
        "container_batch_map": None,
        "hph": rail_row["hph"],
        "loading_line": None,
        "trip_seq": None,
        "reconciled_at": None,
        "reconcile_source_ref": None,
        "dispatch_train_code": None,
    }


def upsert_rows(conn: sqlite3.Connection, table: str, rows: Iterable[dict]) -> int:
    count = 0
    for row in rows:
        cols = list(row.keys())
        conn.execute(
            f"INSERT OR REPLACE INTO {table} ({','.join(cols)}) VALUES ({','.join('?' for _ in cols)})",
            [row[c] for c in cols],
        )
        count += 1
    return count


def refresh_existing_hphs(conn: sqlite3.Connection, ydid_to_hph: dict[str, str], now: str) -> tuple[int, int]:
    c_rows = 0
    for ydid, hph in ydid_to_hph.items():
        cur = conn.execute(
            """
            UPDATE wagon_container_shipments
               SET hph=COALESCE(NULLIF(hph,''), ?),
                   waybill_no=COALESCE(NULLIF(waybill_no,''), ?),
                   updated_at=?
             WHERE project_id=? AND ydid=?
               AND ((hph IS NULL OR hph='') OR (waybill_no IS NULL OR waybill_no=''))
            """,
            (hph, hph, now, PROJECT, ydid),
        )
        c_rows += cur.rowcount
    w_rows = 0
    for ydid, hph in ydid_to_hph.items():
        cur = conn.execute(
            """
            UPDATE wagon_shipments
               SET hph=COALESCE(NULLIF(hph,''), ?),
                   waybill_no=COALESCE(NULLIF(waybill_no,''), ?),
                   updated_at=?
             WHERE project_id=? AND ydid=?
               AND ((hph IS NULL OR hph='') OR (waybill_no IS NULL OR waybill_no=''))
            """,
            (hph, hph, now, PROJECT, ydid),
        )
        w_rows += cur.rowcount
    return c_rows, w_rows


def recount_batch(conn: sqlite3.Connection, batch_id: str, transport_type: str, now: str) -> None:
    if transport_type == "container":
        count_ydid = conn.execute(
            """
            SELECT COUNT(DISTINCT COALESCE(NULLIF(hph,''), ydid))
            FROM wagon_container_shipments
            WHERE batch_id=?
            """,
            (batch_id,),
        ).fetchone()[0]
    else:
        count_ydid = conn.execute(
            """
            SELECT COUNT(DISTINCT COALESCE(NULLIF(hph,''), ydid))
            FROM wagon_shipments
            WHERE batch_id=?
            """,
            (batch_id,),
        ).fetchone()[0]
    conn.execute(
        """
        UPDATE release_batches
           SET batch_count=?,
               actual_wagon_count=?,
               updated_at=?
         WHERE id=?
        """,
        (count_ydid or 0, count_ydid or 0, now, batch_id),
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    by_hph, ydid_to_hph = load_rail_rows()
    now = now_iso_beijing()

    hub = sqlite3.connect(str(SOP_DB))
    hub.row_factory = sqlite3.Row

    print("== 九三历史 95306 回灌分析 ==")
    touched_batches: list[tuple[str, str]] = []
    plan: list[tuple[BatchSpec, str, list[dict]]] = []
    for spec in SPECS:
        batch_id = batch_id_for(hub, spec.ship_name, spec.batch_sequence)
        hphs = load_excel_hphs(spec)
        matched = [by_hph[hph] for hph in hphs if hph in by_hph]
        missing = [hph for hph in hphs if hph not in by_hph]
        print(
            f"{spec.ship_name} {spec.batch_sequence} {spec.transport_type}: "
            f"excel={len(hphs)} matched_95306={len(matched)} missing={len(missing)}"
        )
        if missing:
            print("  sample missing:", ", ".join(missing[:5]))
        plan.append((spec, batch_id, matched))

    c_fill, w_fill = refresh_existing_hphs(hub, ydid_to_hph, now)
    print(f"现有九三明细回填 hph/waybill: container={c_fill} wagon={w_fill}")

    if not args.apply:
        print("\n[干跑] 加 --apply 才会写库。")
        hub.close()
        return 0

    total_container = 0
    total_bulk = 0
    for spec, batch_id, matched in plan:
        if spec.transport_type == "container":
            rows = []
            for rail_row in matched:
                boxes = rail_row.get("boxes") or []
                for pos, box_no in enumerate(boxes, start=1):
                    rows.append(build_container_row(rail_row, box_no, pos, batch_id, spec.ship_name, now))
            n = upsert_rows(hub, "wagon_container_shipments", rows)
            total_container += n
        else:
            rows = [build_bulk_row(rail_row, batch_id, spec.ship_name, now) for rail_row in matched]
            n = upsert_rows(hub, "wagon_shipments", rows)
            total_bulk += n
        touched_batches.append((spec.ship_name, spec.batch_sequence))
        print(f"写入 {spec.ship_name} {spec.batch_sequence}: {n} 行")

    for spec, batch_id, _matched in plan:
        recount_batch(hub, batch_id, spec.transport_type, now)

    hub.commit()
    hub.close()

    for spec, batch_id, _matched in plan:
        compute_for_release_batch(batch_id, db_path=str(SOP_DB))

    print(f"\n完成: container_rows={total_container} bulk_rows={total_bulk}")
    print("涉及批次:", ", ".join(f"{ship}-{lot}" for ship, lot in touched_batches))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
