#!/usr/bin/env python3
"""Generate Jilin Jingang Excel report and factory upload preview"""
import json, sqlite3
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DB = '/Users/qicai21/projects/repos/sop-data-hub/data/sop_agent.db'
OUT_DIR = Path('/Users/qicai21/projects/repos/ops-data-hub/reports')

now = datetime.now()
date_str = now.strftime("%Y%m%d")
time_str = now.strftime("%H%M")

conn = sqlite3.connect(DB)
conn.row_factory = sqlite3.Row
cur = conn.cursor()

# Get departure record
dep = cur.execute("SELECT * FROM departure_records WHERE id LIKE 'dep_jljg_%' ORDER BY created_at DESC LIMIT 1").fetchone()
if not dep:
    print("ERROR: No departure record found")
    exit(1)

wagons = cur.execute("SELECT * FROM wagon_shipments WHERE departure_id=? ORDER BY ticketed_at", (dep['id'],)).fetchall()
wagons = [dict(w) for w in wagons]

filename = f"吉林金钢_发运数据_{date_str}_{time_str}.xlsx"
filepath = OUT_DIR / filename
OUT_DIR.mkdir(parents=True, exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "发运数据"

header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
header_font = Font(bold=True)
header_align = Alignment(horizontal='center', vertical='center')

headers = ['序号', '合同号', '订单标识', '货名', '车号', '箱号', '件数', '装车日期', '进厂日期', '船名']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align

for i, w in enumerate(wagons, 1):
    ws.cell(row=i+1, column=1, value=i)
    ws.cell(row=i+1, column=2, value='JGCG-SFY-HTNK20260501')
    ws.cell(row=i+1, column=3, value='CGR20260518174420')
    ws.cell(row=i+1, column=4, value=w['cargo_name'])
    ws.cell(row=i+1, column=5, value=w['car_no'])
    ws.cell(row=i+1, column=6, value='待补-需从95306API提取箱号')
    ws.cell(row=i+1, column=7, value=1)
    ws.cell(row=i+1, column=8, value=(w.get('ticketed_at') or '')[:10])
    ws.cell(row=i+1, column=9, value='')
    ws.cell(row=i+1, column=10, value='长航滨海')

for col in range(1, len(headers)+1):
    ws.column_dimensions[chr(64+col)].width = 22

wb.save(str(filepath))
print(f"✅ Excel saved: {filepath}")
print(f"   Total rows: {len(wagons)}")
print(f"   ⚠️ Box numbers: placeholder (need extraction from 95306 API)")

# Factory upload preview
print(f"\n=== 工厂上传预览（前3条）===")
for i, w in enumerate(wagons[:3]):
    payload = {
        "wagonNumber": w['car_no'],
        "boxNumber": "待补箱号",
        "goodName": w['cargo_name'],
        "boatName": "长航滨海",
        "contractNumber": "JGCG-SFY-HTNK20260501",
        "orderId": "CGR20260518174420",
        "numberTime": "单次",
        "formId": "MR07",
        "first": "芦家屯装车发运明细"
    }
    print(f"\n--- Car #{i+1}: {w['car_no']} ---")
    print(json.dumps(payload, ensure_ascii=False, indent=2))

conn.close()
