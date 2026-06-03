"""Render a single combined departure excel for tonight's 47 new wagon rows.

跨 lot01/lot02/lot03 三个 release_batch,按 source_message_id 过滤今晚新加的 47 行。
yaml 列定义复用 jilin_jingang.yaml 的 output_templates.departure_excel,但 contract_no /
order_identifier 按每行的 batch_id 现场 lookup,避免单一 ctx。

字段映射(修了 departure_excel.py 的 3 个 key bug):
  - container_no_1 / container_no_2:从 container_numbers_json fallback,优先于 container_no
  - entry_contract_no:从 row.batch → release_batches.contract_no
  - order_identifier:从 row.batch → release_batches.order_identifier
"""

from __future__ import annotations

import json
import sqlite3
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path("/Users/qicai21/projects/repos/sop-data-hub")
SOP_DB = REPO / "data/sop_agent.db"
YAML_PATH = REPO / "config/project_sops/jilin_jingang.yaml"
OUT_DIR = REPO / "output/excel"


def load_template():
    spec = yaml.safe_load(YAML_PATH.read_text(encoding="utf-8"))
    flows = spec.get("flows", {}) or {}
    iters = flows.values() if isinstance(flows, dict) else flows
    for f in iters:
        if not isinstance(f, dict):
            continue
        ot = f.get("output_templates", {})
        if "departure_excel" in ot:
            return ot["departure_excel"]
    raise RuntimeError("departure_excel template not found in yaml")


def fetch_rows(message_inbox_id: str):
    con = sqlite3.connect(str(SOP_DB))
    con.row_factory = sqlite3.Row
    rows = con.execute("""
        SELECT ws.*,
               rb.contract_no AS rb_contract_no,
               rb.order_identifier AS rb_order_identifier,
               rb.cargo_product_name AS rb_cargo_product_name,
               rb.cargo_name AS rb_cargo_name,
               rb.batch_sequence AS rb_lot
        FROM wagon_shipments ws
        LEFT JOIN release_batches rb ON rb.id = ws.batch_id
        WHERE ws.source_message_id = ?
        ORDER BY rb.batch_sequence, ws.car_no, ws.container_no
    """, (message_inbox_id,)).fetchall()
    con.close()
    return [dict(r) for r in rows]


def render(message_inbox_id: str = "229") -> Path:
    tpl = load_template()
    cols = tpl["columns"]
    ws_rows = fetch_rows(message_inbox_id)
    if not ws_rows:
        raise RuntimeError(f"no wagon_shipments rows for source_message_id={message_inbox_id}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tpl["workbook"]["sheet_name"]

    # title
    title = tpl["workbook"]["title"]
    ws.cell(row=title["row"], column=1, value=title["text"])
    if title.get("merge_cells"):
        ws.merge_cells(title["range"])
    ws.cell(row=title["row"], column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=title["row"], column=1).font = Font(bold=True, size=14)

    # headers
    header_row = tpl["workbook"]["header_row"]
    for i, col in enumerate(cols, start=1):
        ws.cell(row=header_row, column=i, value=col["header"])

    # data
    data_start = tpl["workbook"]["data_start_row"]
    for seq, r in enumerate(ws_rows, start=1):
        # container fallback: container_numbers_json > container_no
        containers = []
        if r.get("container_numbers_json"):
            try:
                containers = json.loads(r["container_numbers_json"])
            except Exception:
                containers = []
        if not containers and r.get("container_no"):
            containers = [c.strip() for c in str(r["container_no"]).split("/") if c.strip()]
        c1 = containers[0] if len(containers) >= 1 else ""
        c2 = containers[1] if len(containers) >= 2 else ""

        ta = (r.get("ticketed_at") or "").strip()
        loading_date = ta[:10] if ta and len(ta) >= 10 else ""
        entry_date = ""
        if loading_date:
            try:
                d = datetime.strptime(loading_date, "%Y-%m-%d")
                entry_date = (d + timedelta(days=30)).strftime("%Y-%m-%d")
            except ValueError:
                pass

        # cargo 取顺序:batch.cargo_product_name(真实货物品名,如"印粉")→ wagon.cargo_name
        # (95306 上的品类如"铁矿粉")→ batch.cargo_name(更泛如"铁矿")
        # 业务上 "货物品名" ≠ "货物品类",品名来源是 freight_detail (例:"印粉"),
        # 品类是 95306 / 出港通知单上的大类,优先取品名
        cargo = r.get("rb_cargo_product_name") or r.get("cargo_name") or r.get("rb_cargo_name") or ""

        row_data = {
            "seq": seq,
            "supplier_name": "",
            "detail_account": "",
            "wagon_no": r.get("car_no") or "",
            "container_no_1": c1,
            "container_no_2": c2,
            "loading_date": loading_date,
            "entry_date": entry_date,
            "cargo_name": cargo,
            "ship_name": r.get("ship_name") or "",
            "entry_contract_no": r.get("rb_contract_no") or "",
            "order_identifier": r.get("rb_order_identifier") or "",
            "original_departure_weight": "",
            "shipment_count_type": "单次",
        }

        for ci, col in enumerate(cols, start=1):
            key = col["key"]
            val = col.get("constant")
            if val is None:
                val = row_data.get(key, "")
            ws.cell(row=data_start + seq - 1, column=ci, value=val)

    last_data_row = data_start + len(ws_rows) - 1

    # style: borders
    style = tpl.get("style", {})
    table = style.get("table", {})
    if table.get("bordered_range"):
        side = Side(style=table.get("border_style", "thin"))
        border = Border(top=side, bottom=side, left=side, right=side)
        rng = table["bordered_range"].format(last_data_row=last_data_row)
        for row in ws[rng]:
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

    # column widths
    widths = {1: 6, 2: 18, 3: 12, 4: 10, 5: 14, 6: 14, 7: 12, 8: 12, 9: 14,
              10: 10, 11: 22, 12: 22, 13: 12, 14: 8}
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w

    # filename
    car_count = len(ws_rows)
    fn = tpl["file_naming"]["pattern"].format(
        yyyymmdd=datetime.now().strftime("%Y%m%d"),
        car_count=f"{car_count}行_煤六46节",
    )
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / fn
    wb.save(str(out))
    print(f"saved: {out} ({len(ws_rows)} rows)")
    return out


if __name__ == "__main__":
    mid = sys.argv[1] if len(sys.argv) > 1 else "229"
    render(mid)
